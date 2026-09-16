
import json
import os
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

def save_data(query: str, data: list, export_format: str, logger_func=print):
    # Menyimpan hasil scrape ke format pilihan user dengan styling Excel profesional.
    clean_query = query.strip().replace(" ", "_")
    filename_base = f"{clean_query}_jobs"

    if export_format == "json":
        filename = f"{filename_base}.json"
        with open(filename, "w", encoding="utf-8") as f:
            json.dump(data, f, indent=4, ensure_ascii=False)
        logger_func(f"[UTILS] Berhasil export JSON: {os.path.abspath(filename)}")

    elif export_format == "excel":
        filename = f"{filename_base}.xlsx"
        try:
            df = pd.DataFrame(data)
            
            with pd.ExcelWriter(filename, engine='openpyxl') as writer:
                df.to_excel(writer, index=False, sheet_name='JobStreet Results')
                worksheet = writer.sheets['JobStreet Results']
                
             
                # Header: Dark Steel Blue + Teks Putih Bold
                header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
                header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
                
                # Zebra Striping (Baris Genap): Soft Light Gray/Blue
                zebra_fill = PatternFill(start_color="F2F5F9", end_color="F2F5F9", fill_type="solid")
                data_font = Font(name="Segoe UI", size=10, color="333333")
                
                # Border/Garis Tipis Warna Abu-abu
                thin_border = Border(
                    left=Side(style='thin', color='D9D9D9'),
                    right=Side(style='thin', color='D9D9D9'),
                    top=Side(style='thin', color='D9D9D9'),
                    bottom=Side(style='thin', color='D9D9D9')
                )

                # Alignment
                align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
                align_left_top = Alignment(horizontal="left", vertical="top", wrap_text=True)

                # 1. FORMAT HEADER
                worksheet.row_dimensions[1].height = 28
                for cell in worksheet[1]:
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = align_center
                    cell.border = thin_border

                # 2. FORMAT DATA & ZEBRA STRIPING
                for row_idx, row in enumerate(worksheet.iter_rows(min_row=2, max_row=len(data)+1), start=2):
                    worksheet.row_dimensions[row_idx].height = 22  # Tinggi baris data lebih lega
                    is_even_row = (row_idx % 2 == 0)

                    for col_idx, cell in enumerate(row, start=1):
                        cell.font = data_font
                        cell.border = thin_border
                        
                        # Beri warna selang-seling (zebra striping)
                        if is_even_row:
                            cell.fill = zebra_fill
                        
                        # Kolom Tanggal & Tipe Kerja di-center agar rapi
                        if col_idx in [4, 9]:  
                            cell.alignment = align_center
                        else:
                            cell.alignment = align_left_top

                # 3. AUTO-FIT LEBAR KOLOM
                for col in worksheet.columns:
                    max_len = 0
                    col_letter = col[0].column_letter
                    
                    for cell in col:
                        if cell.value:
                            cell_text = str(cell.value).split('\n')[0]
                            if len(cell_text) > max_len:
                                max_len = len(cell_text)
                    
                    # Batasi lebar kolom max 55 agar deskripsi panjang tidak merusak layout
                    adjusted_width = min(max_len + 4, 55)
                    worksheet.column_dimensions[col_letter].width = max(adjusted_width, 15)

            logger_func(f"[UTILS] Berhasil export Excel (Tampilan Elegan): {os.path.abspath(filename)}")
            
        except ModuleNotFoundError:
            logger_func("[ERROR] Package 'openpyxl' belum terinstal! Jalankan: pip install openpyxl")
        except Exception as e:
            logger_func(f"[ERROR] Gagal menyimpan Excel: {e}")